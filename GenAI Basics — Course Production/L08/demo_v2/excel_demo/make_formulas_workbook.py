# -*- coding: utf-8 -*-
"""Reference workbook for L08 slide 5: the Northstar 2025 book with an "Ad return" sheet built from Excel formulas.

    python make_formulas_workbook.py   ->  northstar_2025_formulas.xlsx (RU sheet name) and northstar_2025_formulas_EN.xlsx

What the demo prompt 2 asks the AI for, done by hand so the slide can show a real Excel formula bar:
SUMIFS over the Sales sheet per month and channel, and revenue / spend in the "per $1" columns.
The Sales sheet is cleaned first (the one exact duplicate removed, the text date 10/05/2025 turned into a date),
so the numbers match demo_v2/answers.md.
"""
import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.chart import LineChart, Reference

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / "northstar_2025.xlsx"

TXT = {
    "ru": dict(out="northstar_2025_formulas.xlsx", sheet="Отдача рекламы",
               head=["Месяц", "Search выручка", "Search расход", "Search $ на $1", "Social выручка", "Social расход", "Social $ на $1"],
               chart="Выручка на $1 рекламы, 2025"),
    "en": dict(out="northstar_2025_formulas_EN.xlsx", sheet="Ad return",
               head=["Month", "Search revenue", "Search spend", "Search $ per $1", "Social revenue", "Social spend", "Social $ per $1"],
               chart="Revenue per $1 of advertising, 2025"),
}


def build(lang):
    tx = TXT[lang]
    wb = openpyxl.load_workbook(SRC)
    ws = wb["Sales"]
    seen, dup = set(), None
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, 1)
        if isinstance(cell.value, str):
            d, m, y = cell.value.split("/")
            cell.value = dt.datetime(int(y), int(m), int(d))
            cell.number_format = "yyyy-mm-dd"
        key = tuple(ws.cell(r, c).value for c in range(1, ws.max_column + 1))
        if key in seen and dup is None:
            dup = r
        seen.add(key)
    if dup:
        ws.delete_rows(dup)
    tbl = ws.tables["Sales"]
    tbl.ref = f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}{ws.max_row}"

    ar = wb.create_sheet(tx["sheet"])
    ar.append(tx["head"])
    for i, m in enumerate(range(1, 13), 2):
        ar.cell(i, 1, dt.datetime(2025, m, 1)).number_format = "mmm yyyy"
        for col, ch, src in ((2, "Search", "G"), (3, "Search", "H"), (5, "Social", "G"), (6, "Social", "H")):
            ar.cell(i, col, f'=SUMIFS(Sales!${src}:${src},Sales!$B:$B,"{ch}",Sales!$A:$A,">="&$A{i},Sales!$A:$A,"<"&EDATE($A{i},1))')
            ar.cell(i, col).number_format = "#,##0"
        ar.cell(i, 4, f"=B{i}/C{i}").number_format = "0.00"
        ar.cell(i, 7, f"=E{i}/F{i}").number_format = "0.00"
    for c, w in zip("ABCDEFG", (12, 16, 15, 16, 16, 15, 16)):
        ar.column_dimensions[c].width = w
    for c in ar[1]:
        c.font = openpyxl.styles.Font(bold=True)
    ch = LineChart()
    ch.title = tx["chart"]
    ch.height, ch.width = 7.5, 17
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorGridlines = None
    ch.legend.position = "r"
    for col in (4, 7):
        ch.add_data(Reference(ar, min_col=col, min_row=1, max_row=13), titles_from_data=True)
    ch.set_categories(Reference(ar, min_col=1, min_row=2, max_row=13))
    ar.add_chart(ch, "I2")
    wb.active = wb.index(ar)
    wb.calculation.fullCalcOnLoad = True
    wb.save(HERE / tx["out"])
    print(tx["out"], "rows in Sales:", ws.max_row - 1)


if __name__ == "__main__":
    for lg in ("ru", "en"):
        build(lg)
