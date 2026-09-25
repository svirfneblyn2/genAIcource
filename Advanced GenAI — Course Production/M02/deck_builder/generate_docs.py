# Documentation Generator for Module 02
# Generates M02_Instructor_Console_Click_Guide_and_Homework.md and .docx
import os
import re
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MODULE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(MODULE_DIR, "M02_Instructor_Console_Click_Guide_and_Homework.md")
DOCX_PATH = os.path.join(MODULE_DIR, "M02_Instructor_Console_Click_Guide_and_Homework.docx")
DEST_DIR = r"G:\My Drive\Advanced GenAI — Course Production\M02 — Agent Frameworks and Orchestration — DONE"

BANNED_SLOP = [
    r"\bdelve\b", r"\btestament\b", r"\btapestry\b", r"\bunleash\b",
    r"\bseamless\b", r"\bgame-changer\b", r"\brevolutionary\b", r"\bpivotal\b",
    r"\bbeacon\b", r"\bcornerstone\b", r"\bfoster\b", r"\bembark\b",
    r"\bfurthermore\b", r"it is important to note", r"in conclusion", r"let us explore"
]

MD_CONTENT = """# Module 02: Instructor Console Click-Guide & Zero-Cost Homework Specification

Course: Advanced Generative AI for Engineers
Module 02: Agent Frameworks and Orchestration: State Machines, Multi-Agent Topologies, and Production Reliability

---

## Part 1: Instructor Console Click-by-Click Guide (15-Minute Tour)

### Pedagogical Goal
Demonstrate to students that production agents are deterministic state machines with verifiable state channels, persistent checkpoints, and human-in-the-loop gateways. You do NOT need expensive cloud clusters or billable vector stores. You use local LangGraph Studio, local Python trace inspectors, or LangSmith to anchor abstract graph theory in concrete execution tools.

### Preparation Before Lecture (Open 4 Browser Tabs / Windows)
1. Tab 1: Presentation deck (`presentation_M02_Agent_Frameworks.html`), opened to Slide 26.
2. Tab 2: LangGraph Studio UI / Local Inspector (`http://localhost:2024` or local trace visualizer).
3. Tab 3: LangSmith Tracing Dashboard (`https://smith.langchain.com/` or Arize Phoenix UI `http://localhost:6006`).
4. Tab 4: Local VS Code editor showing the accounts payable reconciliation graph script (`reconcile_graph.py`).

*Tip: Launch the local LangGraph Studio or Ollama server 10 minutes before class so local warmups do not introduce hesitation on screen.*

---

### Step 1: Graph Compilation & Canvas Overview (3 minutes)

#### Where to look:
1. **LangGraph Studio Canvas (Center Pane)**
   - What to show: The visual graph canvas showing compiled Nodes (`ocr_extraction`, `erp_lookup`, `reconciliation_node`, `post_to_ledger`), Directed Edges, and the Conditional Routing Diamond.
   - What to say:
     > "Notice how the orchestrator visualizes our application. This is not a hidden prompt loop or a black-box chain. It is a formal state automaton. Every rectangle is a Python function that takes a state snapshot and returns a state delta. The diamond is a conditional router that deterministically branches based on state variables."

2. **Left Sidebar: State Schema & Channels**
   - What to show: The defined state schema (`messages: Annotated[list, add_messages]`, `invoice_id: str`, `billed_amount: float`, `variance: float`, `retry_counter: int`).
   - What to say:
     > "This is Layer 1 of our architecture. Every node execution communicates solely through these typed channels. Overwriting is prevented by explicit append reducers."

---

### Step 2: Injecting Test Payload & Live Stepping (3 minutes)

#### Where to click:
1. **Studio Bottom Input Box -> Input JSON**
   - Paste the test payload:
     `{"invoice_id": "INV-8912", "vendor": "Acme Industrial", "billed_amount": 14500.00, "po_number": "PO-4011"}`
   - Click the green button **"Start Run"**.
   - What to show:
     - Watch Node 1 (`ocr_extraction`) illuminate in green as it completes.
     - Watch execution advance to Node 2 (`erp_lookup`).
     - Point out the active thread ID generated in the header (e.g. `thread_rec_9102`).
   - What to say:
     > "Observe the super-step progression. The model does not run unrestricted. The orchestrator executes Node 1, intercepts the state delta, saves an immutable checkpoint to Postgres or SQLite, and then triggers Node 2 to query the ERP sandbox."

---

### Step 3: Triggering the Breakpoint Barrier (3 minutes)

#### Where to look:
1. **Visual Canvas -> Interrupt Barrier before `post_to_ledger`**
   - Point out that execution has halted with status `INTERRUPTED`. Node 4 (`post_to_ledger`) is bordered in yellow/amber and has NOT fired.
   - What to show:
     - Open the **State Inspector Drawer** on the right side.
     - Show the computed state fields:
       - `billed_amount: 14500.00`
       - `po_authorized_amount: 14050.00`
       - `variance: 450.00`
       - `approval_status: "PENDING_CFO_AUTHORIZATION"`
   - What to say:
     > "This is Slide 17 in action. We compiled our graph with interrupt_before=['post_to_ledger']. Because the variance is $450.00, corporate policy mandates human review. Notice that our server process is completely idle. State is persisted in the checkpointer. No CPU is spinning; no tokens are burning. The graph waits safely for an external resume signal."

---

### Step 4: Live State Mutation & Graph Resumption (3 minutes)

#### Where to click:
1. **State Inspector -> Edit Field**
   - Click the field `approved_by` and enter `"cfo_alice@enterprise.com"`.
   - Change `approved_amount` to `14050.00` (demonstrating early payment discount deduction).
   - Click **"Update State"**.
   - What to say:
     > "Notice this crucial architectural point: Human-in-the-Loop is not just a binary yes/no button. Human review is active state mutation. The CFO noticed an unapplied discount and edited the state directly in the checkpointer."

2. **Top Navigation -> Click "Resume"**
   - What to show:
     - The graph unlocks.
     - Node 4 (`post_to_ledger`) illuminates in green.
     - Final output displays: `Journal entry JE-9921 created for $14,050.00. Approved by cfo_alice@enterprise.com.`
   - What to say:
     > "The graph resumed with the sanitized human parameters. We have achieved complete operational safety without breaking automation."

---

### Step 5: Distributed Trace Waterfall Audit (3 minutes)

#### Where to click:
1. **Switch to Tab 3 (LangSmith / Arize Phoenix UI)**
   - Click the latest run matching `thread_rec_9102`.
   - What to show:
     - The complete 4-span waterfall trace.
     - Span 1: `supervisor_plan` (LLM call, 850ms, 1,200 tokens).
     - Span 2: `tool_ocr` (HTTP request, 1,150ms, 0 LLM tokens).
     - Span 3: `erp_lookup` (SQL query, 420ms, 0 LLM tokens).
     - Span 4: `post_to_ledger` (Execution post-resume, 310ms).
     - Total run metrics: Latency 2,730ms active compute, Total Cost: $0.018.
   - What to say:
     > "Here is full observability. Every decision, latency bottleneck, and token cost is accounted for. If a model hallucinates an argument, the trace captures the exact schema rejection before any downstream damage occurs."

---

## Part 2: Student Homework Specification (100% Zero-Cost Guarantee)

### Assignment Title
**Enterprise Agent Engineering Package: Autonomous Accounts Payable & Claims Reconciliation Agent**

### Deadline & Format
- **Format**: 1 Architecture Design Document (PDF or Markdown) OR 1 Local Python Repository with runnable code.
- **Cost Guarantee**: **$0.00 (Zero financial cost guaranteed)**. Students do NOT need cloud subscriptions, credit cards, or paid vector databases.

---

### Business Scenario
You are the Principal AI Engineer for *Apex Enterprise Systems*. The organization processes 25,000 vendor invoices per month. Currently, accounts payable teams manually cross-reference invoices against Purchase Orders in SAP, verify line-item arithmetic, and route discrepancies to finance managers.

Your objective is to design and build a resilient, cyclic, state-machine agent with:
1. **Centralized Typed State**: Explicit schema tracking messages, invoice metadata, line items, retry counts, and audit logs.
2. **Tool Sandboxing**: Isolated tool functions for database lookup and payment ledger posting with parameter validation.
3. **Cyclic Error Recovery**: Automatic reflection loop when database queries fail or parameters fail schema validation (max 3 retries).
4. **Human-in-the-Loop Gateway**: Deterministic breakpoint (`interrupt_before`) triggered whenever discrepancy variance exceeds $0.00 or confidence drops below 95%.
5. **State Mutability**: Capability to resume execution with human-edited approval fields.

---

### Zero-Cost Implementation Options for Students

Students must choose **Option A** (Architecture Track) or **Option B** (Coding Track). Both receive equal grading weight.

#### Option A: Architecture & State Design Specification (No Coding Required)
Ideal for System Architects, Technical Product Managers, and Solution Consultants.
- **Tools**: Mermaid.js, Draw.io, Markdown, or PDF editor. Cost: $0.00.
- **Deliverable Requirements**:
  1. **State Machine Blueprint**: A complete Mermaid graph diagram depicting all Nodes (`ocr_parser`, `erp_match`, `reconcile_evaluator`, `hitl_gate`, `post_ledger`, `fallback_handler`), Static Edges, and Conditional Router Diamonds with explicit branch labels.
  2. **State Schema Specification**: Concrete Python `TypedDict` and `Pydantic` schema definitions including channel types and message reducers (`Annotated[list, add_messages]`).
  3. **Tool Contract & Error Policy Matrix**: Table detailing input JSON schemas, Pydantic field constraints, and specific handling rules across the 4 fault tiers (Transient 5xx, Schema ValidationError, Permanent 401/403, and Dead-Letter Queue).
  4. **HITL Protocol Specification**: Sequence diagram demonstrating state freeze, webhook payload generation, human review interface schema, and resumption payload validation.
  5. **Observability Plan**: Specification of OpenTelemetry span names, attributes (`gen_ai.system`, `gen_ai.usage.input_tokens`), and loop governor threshold limits.

#### Option B: Working Local Python Prototype (Coding Track)
Ideal for Software Engineers, ML Engineers, and Backend Developers.
- **Tools**: Python 3.10+, LangGraph, LangChain Core, and a 100% free local model via Ollama (`llama3.2:3b` or `qwen2.5:7b`) OR Google AI Studio free tier API key. Cost: $0.00.
- **Deliverable Requirements**:
  1. **Source Code**: Clean Python repository containing:
     - `state.py`: TypedDict state definition with `add_messages` reducer.
     - `tools.py`: Mock ERP query tool and ledger posting tool with explicit Pydantic schemas.
     - `graph.py`: StateGraph definition with cyclic retry edge, conditional router, and `MemorySaver` / `SqliteSaver` checkpointer.
     - `main.py`: Executable demonstration script.
  2. **Execution Demonstration**:
     - Run 1 (Auto-Approve Path): Passes matching invoice ($0 variance), runs through to completion without pausing.
     - Run 2 (HITL Path): Passes discrepancy invoice ($450 variance), hits `interrupt_before`, demonstrates state inspection, applies `app.update_state()` to simulate approval, and resumes to completion.
  3. **README.md**: Setup instructions to run locally in under 3 minutes with zero API keys required (using Ollama).

---

### 100-Point Grading Rubric

| Criterion | Points | Excellent (Full Points) | Needs Improvement (Partial) | Unacceptable (Zero) |
| :--- | :--- | :--- | :--- | :--- |
| **1. State Schema Architecture** | **25 pts** | Explicit TypedDict/Pydantic schema with typed channels, append-only message reducers, and audit fields. | Generic untyped dictionary used; lacks explicit reducers or type hints. | Missing state schema; global variables or uncontrolled state. |
| **2. Cyclic Flow & Error Reflection** | **25 pts** | Cyclic graph edges enable self-correction; retry counter enforces hard cap (max 3); error payloads feed back to model. | Retries implemented as simple while-loop outside graph; lacks reflection feedback. | Forward-only DAG with no error recovery; crashes on failure. |
| **3. Tool Sandboxing & Contracts** | **20 pts** | Explicit JSON/Pydantic schemas with type constraints; tools wrapped in safe execution boundary with timeout. | Tools lack validation schemas; unhandled exceptions bubble to main process. | Model generates raw unparsed code with direct host execution. |
| **4. HITL Gateway & Mutability** | **15 pts** | Deterministic breakpoint before destructive action; state saved in checkpointer; demonstrates state mutation on resume. | Breakpoint pauses but cannot accept mutated state; resume logic is binary only. | No human-in-the-loop protection; autonomous execution of sensitive action. |
| **5. Observability & Governance** | **15 pts** | Recursion limits configured; cost/token governors defined; trace spans capture latency and token consumption. | Basic logging present but lacks span hierarchy or loop circuit breakers. | Zero telemetry, zero logging, and unbounded recursion depth. |
| **TOTAL** | **100 pts** | **Minimum Passing Score: 80 / 100** | | |

---

### Submission Guidelines
1. Package your work as a single zip archive or GitHub repository link.
2. Ensure all diagrams and code include zero proprietary corporate data.
3. Include a 2-minute screen recording (Loom or MP4) or a comprehensive markdown walkthrough verifying the execution paths.
"""

def generate_docx():
    doc = Document()
    
    # Page setup: Margins 0.8 inch
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Styles & Fonts
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1a, 0x20, 0x2c) # Slate

    # Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("Module 02: Instructor Console Click-Guide & Zero-Cost Homework")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
    p_title.paragraph_format.space_after = Pt(4)

    # Subtitle
    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run("Course: Advanced Generative AI for Engineers | Module 02: Agent Frameworks & Orchestration")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x3b, 0x82, 0xf6)
    p_sub.paragraph_format.space_after = Pt(16)

    # Parse Markdown lines into docx
    lines = MD_CONTENT.split("\n")
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    for line in lines:
        stripped = line.strip()
        
        # Code block handling
        if stripped.startswith("```"):
            if in_code_block:
                # Flush code block
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.3)
                p_code.paragraph_format.space_before = Pt(4)
                p_code.paragraph_format.space_after = Pt(4)
                run_c = p_code.add_run("\n".join(code_lines))
                run_c.font.name = 'Consolas'
                run_c.font.size = Pt(9.5)
                run_c.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Table handling
        if stripped.startswith("|"):
            if "---" in stripped:
                continue # delimiter row
            cols = [c.strip() for c in stripped.split("|")[1:-1]]
            if not in_table:
                in_table = True
                table_rows = [cols]
            else:
                table_rows.append(cols)
            continue
        elif in_table:
            # Flush table
            if table_rows:
                num_cols = len(table_rows[0])
                t = doc.add_table(rows=len(table_rows), cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_data in enumerate(table_rows):
                    for c_idx, cell_data in enumerate(row_data):
                        cell = t.cell(r_idx, c_idx)
                        clean_text = cell_data.replace("**", "").strip()
                        cell.text = clean_text
                        p_cell = cell.paragraphs[0]
                        p_cell.paragraph_format.space_before = Pt(2)
                        p_cell.paragraph_format.space_after = Pt(2)
                        run_cell = p_cell.runs[0] if p_cell.runs else p_cell.add_run(clean_text)
                        run_cell.font.name = 'Calibri'
                        run_cell.font.size = Pt(9.5)
                        if r_idx == 0:
                            run_cell.font.bold = True
                            run_cell.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
            in_table = False
            table_rows = []

        # Headings
        if stripped.startswith("## "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            r = h.add_run(stripped[3:].replace("**", ""))
            r.font.name = 'Calibri'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
        elif stripped.startswith("### "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            r = h.add_run(stripped[4:].replace("**", ""))
            r.font.name = 'Calibri'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            r = h.add_run(stripped[5:].replace("**", ""))
            r.font.name = 'Calibri'
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif stripped.startswith("> "):
            # Callout quote
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Inches(0.25)
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(4)
            r_q = p_q.add_run(stripped[2:])
            r_q.font.italic = True
            r_q.font.color.rgb = RGBColor(0x25, 0x63, 0xeb)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p_bullet = doc.add_paragraph(style='List Bullet')
            p_bullet.paragraph_format.space_before = Pt(2)
            p_bullet.paragraph_format.space_after = Pt(2)
            r_b = p_bullet.add_run(stripped[2:].replace("**", ""))
            r_b.font.size = Pt(10.5)
        elif re.match(r"^\d+\.\s", stripped):
            p_num = doc.add_paragraph(style='List Number')
            p_num.paragraph_format.space_before = Pt(2)
            p_num.paragraph_format.space_after = Pt(2)
            idx_space = stripped.find(" ")
            r_n = p_num.add_run(stripped[idx_space+1:].replace("**", ""))
            r_n.font.size = Pt(10.5)
        elif stripped == "---":
            continue
        elif stripped:
            p_norm = doc.add_paragraph()
            p_norm.paragraph_format.space_before = Pt(3)
            p_norm.paragraph_format.space_after = Pt(3)
            r_norm = p_norm.add_run(stripped.replace("**", ""))
            r_norm.font.size = Pt(10.5)

    doc.save(DOCX_PATH)
    print(f"  [SUCCESS] Wrote Word DOCX deliverable to: {DOCX_PATH}")

def build_docs():
    print("================================================================================")
    print("STARTING MODULE 02 DOCUMENTATION GENERATION PIPELINE")
    print("================================================================================")

    # Step 1: Validate Markdown for zero em-dashes and en-dashes
    print("Step 1: Validating editorial punctuation integrity...")
    em_dash_matches = [m.start() for m in re.finditer(r"\u2014|—|&mdash;|&#8212;", MD_CONTENT)]
    en_dash_matches = [m.start() for m in re.finditer(r"\u2013|–|&ndash;|&#8211;", MD_CONTENT)]

    if em_dash_matches:
        for pos in em_dash_matches[:5]:
            snippet = MD_CONTENT[max(0, pos-40):min(len(MD_CONTENT), pos+40)]
            print(f"  [ERROR] Em-dash found at pos {pos}: ...{snippet}...")
        raise ValueError(f"Punctuation check failed: {len(em_dash_matches)} em-dashes found!")

    if en_dash_matches:
        for pos in en_dash_matches[:5]:
            snippet = MD_CONTENT[max(0, pos-40):min(len(MD_CONTENT), pos+40)]
            print(f"  [ERROR] En-dash found at pos {pos}: ...{snippet}...")
        raise ValueError(f"Punctuation check failed: {len(en_dash_matches)} en-dashes found!")

    print("  [SUCCESS] Zero em-dashes and zero en-dashes confirmed.")

    # Step 2: Anti-slop vocabulary check
    print("Step 2: Validating anti-slop vocabulary rules...")
    for pattern in BANNED_SLOP:
        matches = list(re.finditer(pattern, MD_CONTENT, re.IGNORECASE))
        if matches:
            for m in matches[:3]:
                snippet = MD_CONTENT[max(0, m.start()-30):min(len(MD_CONTENT), m.end()+30)]
                print(f"  [ERROR] Banned slop found: '{pattern}' -> ...{snippet}...")
            raise ValueError(f"Anti-slop check failed for pattern '{pattern}'!")

    print("  [SUCCESS] Anti-slop check passed cleanly.")

    # Step 3: Write Markdown
    print("Step 3: Writing Markdown deliverable...")
    with open(MD_PATH, "w", encoding="utf-8") as f:
        f.write(MD_CONTENT)
    print(f"  [SUCCESS] Wrote Markdown deliverable to: {MD_PATH} ({len(MD_CONTENT)} bytes)")

    # Step 4: Write Word DOCX
    print("Step 4: Compiling Word DOCX deliverable...")
    generate_docx()

    # Step 5: Deploy to Google Drive Production Folder
    print(f"Step 5: Copying deliverables to Google Drive folder:\n  {DEST_DIR}")
    if os.path.exists(DEST_DIR):
        dest_md = os.path.join(DEST_DIR, os.path.basename(MD_PATH))
        dest_docx = os.path.join(DEST_DIR, os.path.basename(DOCX_PATH))
        shutil.copy2(MD_PATH, dest_md)
        shutil.copy2(DOCX_PATH, dest_docx)
        print(f"  [SUCCESS] Copied Markdown: {dest_md} ({os.path.getsize(dest_md)} bytes)")
        print(f"  [SUCCESS] Copied DOCX:     {dest_docx} ({os.path.getsize(dest_docx)} bytes)")
    else:
        print(f"  [WARNING] Destination directory {DEST_DIR} does not exist. Skipping copy.")

    print("================================================================================")
    print("MODULE 02 DOCUMENTATION PIPELINE COMPLETE: ALL CRITERIA SATISFIED!")
    print("================================================================================")

if __name__ == "__main__":
    build_docs()
